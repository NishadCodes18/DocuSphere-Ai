from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pptx import Presentation
from pypdf import PdfReader


@dataclass
class SourceBlock:
    text: str
    page_number: int | None = None
    slide_number: int | None = None


def parse_file(path: str, media_type: str) -> list[SourceBlock]:
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf" or media_type == "application/pdf":
        try:
            reader = PdfReader(path)
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception:
                    raise ValueError("This PDF file is password protected.")
            blocks: list[SourceBlock] = []
            for idx, page in enumerate(reader.pages, start=1):
                try:
                    text = (page.extract_text() or "").strip()
                except Exception:
                    text = ""
                if text:
                    blocks.append(SourceBlock(text=text, page_number=idx))
            return blocks
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"Unable to read PDF file: {e}")

    if suffix == ".docx":
        try:
            doc = DocxDocument(path)
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip()).strip()
            return [SourceBlock(text=text)] if text else []
        except Exception as e:
            raise ValueError(f"Unable to read DOCX file: {e}")

    if suffix == ".pptx":
        try:
            prs = Presentation(path)
            blocks = []
            for slide_no, slide in enumerate(prs.slides, start=1):
                parts: list[str] = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        parts.append(shape.text.strip())
                text = "\n".join(parts).strip()
                if text:
                    blocks.append(SourceBlock(text=text, slide_number=slide_no))
            return blocks
        except Exception as e:
            raise ValueError(f"Unable to read PPTX file: {e}")

    if suffix in (".txt", ".md", ".csv") or media_type.startswith("text/"):
        try:
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read().strip()
            return [SourceBlock(text=text, page_number=1)] if text else []
        except Exception as e:
            raise ValueError(f"Unable to read text file: {e}")

    raise ValueError("Unsupported file type. Use PDF, DOCX, PPTX, TXT, or MD.")


def chunk_blocks(blocks: list[SourceBlock], chunk_size: int = 900, overlap: int = 120) -> list[SourceBlock]:
    if overlap >= chunk_size:
        raise ValueError("overlap must be smaller than chunk_size")

    chunks: list[SourceBlock] = []
    for block in blocks:
        words = block.text.split()
        start = 0
        while start < len(words):
            end = min(start + chunk_size, len(words))
            content = " ".join(words[start:end]).strip()
            if content:
                chunks.append(SourceBlock(content, block.page_number, block.slide_number))
            if end == len(words):
                break
            start = end - overlap
    return chunks
